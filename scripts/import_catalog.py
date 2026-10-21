#!/usr/bin/env python3
"""
Loads a supermarket catalog spreadsheet into the app's catalog_products table (Supabase).
Built for the .xlsx that mercadona_bcn_catalog.py writes; other stores can use the same columns.

Setup (once):  pip install requests openpyxl
  export SUPABASE_URL="https://<your-project-ref>.supabase.co"
  export SUPABASE_SERVICE_KEY="<service_role key>"   # Supabase > Project Settings > API. SECRET: never commit or share it.

Run:
  python scripts/import_catalog.py mercadona_bcn_catalog.xlsx --store mercadona --dry-run   # check, writes nothing
  python scripts/import_catalog.py mercadona_bcn_catalog.xlsx --store mercadona             # add new + update existing
  python scripts/import_catalog.py mercadona_bcn_catalog.xlsx --store mercadona --replace   # remove the store's old rows first

--store is the id used in STORES in lib/catalog.ts (lower case, e.g. mercadona, carrefour).
"""
import argparse
import json
import os
import sys

import requests
from openpyxl import load_workbook

# Spreadsheet header -> what we call it. Matching ignores case and extra spaces, so column order doesn't matter.
HEADERS = {
    "id": "ID", "ean": "EAN", "name": "Product", "brand": "Brand", "category": "Category", "subcategory": "Subcategory",
    "packaging": "Packaging", "price": "Price (€)", "size": "Unit size", "size_unit": "Size unit",
    "kcal": "kcal /100", "protein": "Protein g /100", "carbs": "Carbs g /100", "sugar": "of which sugars g",
    "fat": "Fat g /100", "sat_fat": "of which saturates g", "fiber": "Fibre g /100", "salt": "Salt g /100",
    "image": "Image URL", "url": "Product URL",
}
# store unit -> (unit the nutrition is "per 100" of, how many of that unit one store unit is)
UNITS = {"l": ("ml", 1000), "ml": ("ml", 1), "cl": ("ml", 10), "kg": ("g", 1000), "g": ("g", 1), "gr": ("g", 1)}


def text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = str(v).strip()
    return t or None


def number(v):
    if v is None or v == "":
        return None
    try:
        x = float(str(v).strip().replace(",", ".")) if not isinstance(v, (int, float)) else float(v)
    except ValueError:
        return None
    return round(x, 3) if x >= 0 else None


def convert(store, get):
    """One spreadsheet row -> one catalog_products row, or None if it has no id or name."""
    pid, name = text(get("id")), text(get("name"))
    if not pid or not name:
        return None
    size, size_unit = number(get("size")), (text(get("size_unit")) or "").lower()
    unit, factor = UNITS.get(size_unit, ("g", None))  # items sold by the piece ("ud") have no weight
    image = text(get("image"))
    return {
        "store": store, "product_id": pid, "ean": text(get("ean")), "name": name, "brand": text(get("brand")),
        "category": text(get("category")), "subcategory": text(get("subcategory")), "packaging": text(get("packaging")),
        "price": number(get("price")), "pack_size": size, "pack_unit": size_unit or None, "unit": unit,
        "pack_amount": round(size * factor, 3) if size and factor else None,
        "kcal": number(get("kcal")), "protein": number(get("protein")), "carbs": number(get("carbs")),
        "sugar": number(get("sugar")), "fat": number(get("fat")), "sat_fat": number(get("sat_fat")),
        "fiber": number(get("fiber")), "salt": number(get("salt")),
        "image_url": image.split("?")[0] if image else None,  # the app asks for the size it needs
        "product_url": text(get("url")),
    }


def read_rows(path, store):
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    col = {}
    for key, title in HEADERS.items():
        if title.lower() in header:
            col[key] = header.index(title.lower())
    for need in ("id", "name"):
        if need not in col:
            sys.exit(f'Column "{HEADERS[need]}" not found in {path}. Found: {[h for h in header if h]}')
    out, seen = [], set()
    for r in rows:
        item = convert(store, lambda k: r[col[k]] if k in col and col[k] < len(r) else None)
        if item and item["product_id"] not in seen:
            seen.add(item["product_id"])
            out.append(item)
    return out


def upload(rows, store, replace):
    url, key = os.environ.get("SUPABASE_URL"), os.environ.get("SUPABASE_SERVICE_KEY")
    if not url or not key:
        sys.exit("Set SUPABASE_URL and SUPABASE_SERVICE_KEY first (see the top of this file).")
    base = url.rstrip("/") + "/rest/v1/catalog_products"
    head = {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json"}
    if replace:
        r = requests.delete(base, params={"store": f"eq.{store}"}, headers=head, timeout=60)
        if not r.ok:
            sys.exit(f"Could not clear old {store} rows: {r.status_code} {r.text[:300]}")
        print(f"Cleared old {store} rows")
    for i in range(0, len(rows), 500):
        chunk = rows[i:i + 500]
        r = requests.post(base, params={"on_conflict": "store,product_id"}, data=json.dumps(chunk), timeout=120,
                          headers={**head, "Prefer": "resolution=merge-duplicates,return=minimal"})
        if not r.ok:
            sys.exit(f"Upload failed at row {i}: {r.status_code} {r.text[:300]}")
        print(f"Uploaded {min(i + 500, len(rows))}/{len(rows)}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--store", required=True, help="store id, e.g. mercadona")
    ap.add_argument("--dry-run", action="store_true", help="read and check the file, upload nothing")
    ap.add_argument("--replace", action="store_true", help="delete this store's existing rows before uploading")
    a = ap.parse_args()
    store = a.store.strip().lower()
    data = read_rows(a.xlsx, store)
    no_macros = sum(1 for d in data if d["kcal"] is None)
    print(f"{len(data)} products read, {no_macros} without calories (they show as 'No nutrition data')")
    if a.dry_run:
        for d in data[:3]:
            print(json.dumps(d, ensure_ascii=False))
    else:
        upload(data, store, a.replace)
        print("Done")
